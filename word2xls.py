#!/usr/bin/env python
# -*- coding: utf-8 -*-

import datetime
import glob

from docx import Document
import xlwt.Workbook

TIMESTAMP_FORMAT = '%Y-%m-%dT%H-%M-%S'

algn1 = xlwt.Alignment()
algn1.wrap = 1
style1 = xlwt.XFStyle()
style1.alignment = algn1

docx_files = glob.glob("*.docx")
workbook = xlwt.Workbook()
file_count = 0

for docfilename in docx_files:
    file_count += 1
    document = Document(docfilename)

    curriculum_changes = document.tables[1]

    sheet = workbook.add_sheet(docfilename)
    row_count = 0
    for r in curriculum_changes.rows:
        cell_count = 0
        for c in r.cells:
            sheet.write(row_count, cell_count, c.text, style1)
            #print "row %s, cell %s = [%s]" % (row_count, cell_count, c.text)
            cell_count += 1
        row_count += 1

filename = 'docs_%s.xls' % (datetime.datetime.strftime(datetime.datetime.utcnow(), TIMESTAMP_FORMAT))
workbook.save(filename)
print "%s docx file(s) converted to %s" % (file_count, filename)
